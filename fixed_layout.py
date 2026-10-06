"""
Фиксированная вёрстка: страница PDF -> страница DOCX, на которой всё стоит там же, где в PDF.

- Всё нетекстовое (рамка документа, основная надпись, линии, картинки, повёрнутый текст на полях) —
  картинка-подложка за текстом. Она привязана к странице и сдвинуться не может.
- Таблицы с линиями — настоящие таблицы Word, закреплённые на странице: ширины колонок и высоты строк
  как в PDF, объединённые ячейки, толщина рамок ячеек, заливка, текст в своих ячейках.
- Текст — абзацы в рамках (framePr), закреплённые на странице: строки разбиты как в оригинале,
  шрифт, кегль, жирность, курсив и цвет сохраняются, текст редактируется.

Рамки документа и штампы (таблица, у которой есть ячейка больше четверти страницы) таблицами Word не
делаются: их линии остаются на подложке, а надписи в них — текстом на своих местах.
"""
import bisect
import io
import re
import statistics

import pymupdf
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

if hasattr(pymupdf, "no_recommend_layout"):
    pymupdf.no_recommend_layout()  # иначе find_tables печатает рекламу pymupdf_layout поверх прогресс-бара

TW = 20                 # twips в пункте
BACKGROUND_DPI = 200    # разрешение подложки (рамки, линии, картинки)
SNAP = 1.5              # допуск при сравнении координат линий и ячеек, pt
MAX_PAGE = 1584          # наибольший размер страницы в Word (22 дюйма), pt
INVALID_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿\ud800-\udfff]")
SUBSET = re.compile(r"^[A-Z]{6}\+")
# имена шрифтов PDF -> шрифты Windows (свободные аналоги — на оригиналы)
FONT_MAP = {
    "timesnewroman": "Times New Roman", "times": "Times New Roman", "liberationserif": "Times New Roman",
    "tinos": "Times New Roman", "nimbusroman": "Times New Roman", "tiro": "Times New Roman",
    "arial": "Arial", "liberationsans": "Arial", "arimo": "Arial", "helvetica": "Arial", "helv": "Arial",
    "nimbussans": "Arial", "couriernew": "Courier New", "courier": "Courier New", "cour": "Courier New",
    "liberationmono": "Courier New", "cousine": "Courier New", "calibri": "Calibri", "cambria": "Cambria",
    "verdana": "Verdana", "tahoma": "Tahoma", "georgia": "Georgia", "segoeui": "Segoe UI",
    "gosttypea": "GOST type A", "gosttypeb": "GOST type B", "isocpeur": "ISOCPEUR", "isocp": "ISOCPEUR",
    "symbol": "Symbol", "wingdings": "Wingdings", "ptserif": "PT Serif", "ptsans": "PT Sans",
}
SYMBOL_BULLETS = {"": "•", "": "▪", "": "➢", "": "❖", "": "✓",
                  "": "➔", "": "■", "": "❑"}


# ----------------------------------------------------------------------------- текст страницы

def word_font(pdf_font):
    """Имя шрифта PDF ('ABCDEF+TimesNewRomanPS-BoldMT') -> имя шрифта Word ('Times New Roman')."""
    name = SUBSET.sub("", pdf_font or "")
    base = re.sub(r"(PSMT|PSM|PS|MT)$", "", re.split(r"[-,]", name)[0])
    key = re.sub(r"[^a-z0-9]", "", base.lower())
    for prefix in sorted(FONT_MAP, key=len, reverse=True):
        if key.startswith(prefix):
            return FONT_MAP[prefix]
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", base) or "Times New Roman"


def span_style(span):
    """Оформление куска текста для Word."""
    name = (span.get("font") or "").lower()
    font = word_font(span.get("font"))
    return {
        "font": font,
        "size": round(span["size"] * 2) / 2,
        "bold": bool(span["flags"] & 16) or any(w in name for w in ("bold", "black", "heavy", "semibold")),
        "italic": bool(span["flags"] & 2) or "italic" in name or "oblique" in name,
        "color": f"{span.get('color', 0) & 0xFFFFFF:06X}",
        "super": bool(span["flags"] & 1),
        "symbol": font in ("Symbol", "Wingdings"),
    }


def text_items(page):
    """Куски текста страницы: строки, разрезанные по большим промежуткам (графы, колонки, «подпись ___ дата»).
    У каждого куска — рамка, базовая линия, кегль, направление и «прогоны» одинаково оформленного текста."""
    items = []
    flags = pymupdf.TEXT_PRESERVE_WHITESPACE | pymupdf.TEXT_PRESERVE_LIGATURES | pymupdf.TEXT_MEDIABOX_CLIP
    for block_no, block in enumerate(page.get_text("rawdict", flags=flags)["blocks"]):
        for line in block.get("lines", []):
            dx, dy = line["dir"]
            horizontal = abs(dy) < 0.05 and dx > 0
            chunk, last_x1 = None, None
            for span in line["spans"]:
                if span.get("alpha", 255) == 0:
                    continue
                style = span_style(span)
                for char in span["chars"]:
                    c = INVALID_XML.sub("", char["c"])
                    if not c:
                        continue
                    if style["symbol"] and c in SYMBOL_BULLETS:
                        c, style = SYMBOL_BULLETS[c], dict(style, font="Arial", symbol=False)
                    x0, y0, x1, y1 = char["bbox"]
                    gap_limit = max(1.6 * span["size"], 8)
                    if horizontal and chunk and not c.isspace() and last_x1 is not None and x0 - last_x1 > gap_limit:
                        items.append(chunk)
                        chunk = None
                    if chunk is None:
                        if c.isspace():
                            continue
                        chunk = {"runs": [], "x0": x0, "y0": y0, "x1": x1, "y1": y1, "block": block_no,
                                 "base": char["origin"][1], "size": span["size"], "dir": (dx, dy),
                                 "asc": span.get("ascender", 0.9), "desc": span.get("descender", -0.2)}
                    if chunk["runs"] and chunk["runs"][-1][1] == style:
                        chunk["runs"][-1][0] += c
                    else:
                        chunk["runs"].append([c, style])
                    if not c.isspace():
                        chunk["x0"], chunk["y0"] = min(chunk["x0"], x0), min(chunk["y0"], y0)
                        chunk["x1"], chunk["y1"] = max(chunk["x1"], x1), max(chunk["y1"], y1)
                        chunk["size"] = max(chunk["size"], span["size"])
                        last_x1 = x1
            if chunk:
                items.append(chunk)
    for item in items:  # пробелы по краям куска не нужны
        item["runs"][-1][0] = item["runs"][-1][0].rstrip()
        item["runs"] = [run for run in item["runs"] if run[0]]
    return [item for item in items if item["runs"]]


def is_horizontal(item):
    dx, dy = item["dir"]
    return abs(dy) < 0.05 and dx > 0


# ----------------------------------------------------------------------------- линии и таблицы

def segments(page):
    """Отрезки линий страницы: горизонтальные (y, x0, x1, толщина) и вертикальные (x, y0, y1, толщина),
    а также заливки (прямоугольник, цвет)."""
    horizontal, vertical, fills = [], [], []
    for path in page.get_drawings():
        width = path.get("width") or 0
        stroke = path.get("color") is not None and path.get("type") in ("s", "fs")
        fill = path.get("fill") if path.get("type") in ("f", "fs") else None
        for item in path["items"]:
            if item[0] == "l":
                (x0, y0), (x1, y1) = item[1], item[2]
                if abs(y0 - y1) <= 0.5:
                    horizontal.append((y0, min(x0, x1), max(x0, x1), width or 0.5))
                elif abs(x0 - x1) <= 0.5:
                    vertical.append((x0, min(y0, y1), max(y0, y1), width or 0.5))
            elif item[0] in ("re", "qu"):
                rect = item[1] if item[0] == "re" else item[1].rect
                if fill is not None and not stroke and rect.height <= 2.5 and rect.width > 2.5:
                    horizontal.append(((rect.y0 + rect.y1) / 2, rect.x0, rect.x1, rect.height))  # линия-заливка
                elif fill is not None and not stroke and rect.width <= 2.5 and rect.height > 2.5:
                    vertical.append(((rect.x0 + rect.x1) / 2, rect.y0, rect.y1, rect.width))
                else:
                    if stroke:
                        horizontal += [(rect.y0, rect.x0, rect.x1, width), (rect.y1, rect.x0, rect.x1, width)]
                        vertical += [(rect.x0, rect.y0, rect.y1, width), (rect.x1, rect.y0, rect.y1, width)]
                    if fill is not None:
                        fills.append((rect, fill))
    return horizontal, vertical, fills


def edge_width(lines, at, start, end):
    """Толщина линии, проходящей по краю ячейки (0 — линии нет)."""
    covered, width = [], 0.0
    for pos, a, b, w in lines:
        if abs(pos - at) <= SNAP + w / 2 and b > start and a < end:
            covered.append((max(a, start), min(b, end)))
            width = max(width, w)
    covered.sort()
    total, reach = 0.0, start
    for a, b in covered:
        if b > reach:
            total += b - max(a, reach)
            reach = b
    return width if end > start and total >= 0.6 * (end - start) else 0.0


def cluster(values):
    """Близкие координаты (в пределах SNAP) — одна линия сетки."""
    result = []
    for v in sorted(values):
        if result and v - result[-1][-1] <= SNAP:
            result[-1].append(v)
        else:
            result.append([v])
    return [sum(group) / len(group) for group in result]


def nearest(grid, value):
    return min(range(len(grid)), key=lambda i: abs(grid[i] - value))


def find_word_tables(page):
    """Таблицы с линиями, которые станут таблицами Word. Рамка документа со штампом (ячейка больше
    четверти страницы) и одиночные прямоугольники таблицами не считаются."""
    area = abs(page.rect)
    try:
        found = page.find_tables(strategy="lines").tables
    except Exception:  # noqa: BLE001 — нестандартная графика: тогда всё уйдёт на подложку
        return []
    tables = []
    for table in found:
        cells = [pymupdf.Rect(c) for c in table.cells if c]
        if len(cells) < 2 or any(abs(c) > 0.25 * area for c in cells) or abs(pymupdf.Rect(table.bbox)) > 0.85 * area:
            continue
        xs = cluster([c.x0 for c in cells] + [c.x1 for c in cells])
        ys = cluster([c.y0 for c in cells] + [c.y1 for c in cells])
        if len(xs) < 2 or len(ys) < 2:
            continue
        grid = []
        for c in cells:
            c0, c1, r0, r1 = nearest(xs, c.x0), nearest(xs, c.x1), nearest(ys, c.y0), nearest(ys, c.y1)
            if c1 > c0 and r1 > r0:
                grid.append({"rect": c, "c0": c0, "c1": c1, "r0": r0, "r1": r1, "items": []})
        if len(grid) >= 2:
            tables.append({"bbox": pymupdf.Rect(xs[0], ys[0], xs[-1], ys[-1]), "xs": xs, "ys": ys, "cells": grid})
    return tables


# ----------------------------------------------------------------------------- XML для Word

def twips(pt):
    return int(round(pt * TW))


def run_xml(text, style):
    color = "" if style["color"] == "000000" else f'<w:color w:val="{style["color"]}"/>'
    font = style["font"].replace('"', "")
    props = (f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}" w:eastAsia="{font}"/>'
             + ("<w:b/><w:bCs/>" if style["bold"] else "") + ("<w:i/><w:iCs/>" if style["italic"] else "")
             + color + f'<w:sz w:val="{int(style["size"] * 2)}"/><w:szCs w:val="{int(style["size"] * 2)}"/>'
             + ('<w:vertAlign w:val="superscript"/>' if style["super"] else ""))
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return f'<w:r><w:rPr>{props}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'


def lines_xml(lines):
    """Строки (каждая — список прогонов) через разрыв строки, как в оригинале."""
    parts = []
    for k, runs in enumerate(lines):
        if k:
            parts.append("<w:r><w:br/></w:r>")
        parts += [run_xml(text, style) for text, style in runs]
    return "".join(parts)


def line_metrics(item, pitch=None):
    """(высота строки, расстояние от верха строки до базовой линии) в pt — так строка в Word встаёт
    на ту же базовую линию, что и в PDF (при точном межстрочном интервале лишнее место — сверху)."""
    descent = -item["desc"] * item["size"]
    natural = (item["asc"] - item["desc"]) * item["size"]
    height = pitch if pitch and pitch > 0.6 * natural else natural
    return height, height - descent


def frame_paragraph(x, y, width, height, line, align, first_indent, content):
    return parse_xml(
        f'<w:p {nsdecls("w")}><w:pPr>'
        f'<w:framePr w:w="{twips(width)}" w:h="{twips(height)}" w:hRule="atLeast" w:hSpace="0" w:vSpace="0" '
        f'w:wrap="around" w:vAnchor="page" w:hAnchor="page" w:x="{twips(x)}" w:y="{twips(y)}"/>'
        f'<w:widowControl w:val="0"/><w:autoSpaceDE w:val="0"/><w:autoSpaceDN w:val="0"/><w:snapToGrid w:val="0"/>'
        f'<w:spacing w:before="0" w:after="0" w:line="{twips(line)}" w:lineRule="exact"/>'
        f'<w:ind w:left="0" w:right="0" w:firstLine="{twips(first_indent)}"/><w:jc w:val="{align}"/>'
        f'</w:pPr>{content}</w:p>')


def spacer_paragraph():
    """Пустой абзац минимальной высоты (Word требует абзац после таблицы и в конце страницы)."""
    return parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:before="0" w:after="0" w:line="20" '
                     f'w:lineRule="exact"/></w:pPr><w:r><w:rPr><w:sz w:val="2"/></w:rPr></w:r></w:p>')


# ----------------------------------------------------------------------------- абзацы-рамки

def group_items(items):
    """Строки, идущие подряд с одинаковым шагом, кеглем и левым краем, — в один абзац (строки остаются
    строками, через разрыв строки). Строки разных текстовых блоков объединяются, только если их левые края
    совпадают (у распознанных сканов PyMuPDF иногда делит абзац на несколько блоков)."""
    groups = []
    for item in sorted(items, key=lambda it: (it["block"], it["base"], it["x0"])):
        group = groups[-1] if groups else None
        if group and continues(group, item, item["block"] == group[-1]["block"]):
            group.append(item)
        else:
            groups.append([item])
    return groups


def continues(group, item, same_block):
    first, last, size = group[0], group[-1], group[-1]["size"]
    step = item["base"] - last["base"]
    if abs(item["size"] - size) > 0.15 * size or not 0.8 * size <= step <= 2.2 * size:
        return False
    if not same_block and (abs(item["size"] - size) > 0.5 or step > 1.6 * size or abs(item["x0"] - first["x0"]) > 2
                           and (len(group) == 1 or abs(item["x0"] - group[1]["x0"]) > 2)):
        return False  # строка другого блока — только точное продолжение абзаца, а не заголовок + текст
    if len(group) > 1:
        pitch = group[1]["base"] - first["base"]
        return abs(step - pitch) <= 0.12 * pitch and abs(item["x0"] - group[1]["x0"]) <= 2
    indent = first["x0"] - item["x0"]  # первая строка может начинаться с красной строки
    return -2 <= indent <= 2 or indent <= max(4.5 * size, 45) and first["x1"] >= item["x1"] - 2


def right_edges(items):
    """Правые края колонок текста: x, на котором кончаются хотя бы три строки страницы."""
    ends = sorted(it["x1"] for it in items)
    return [x for x in ends if bisect.bisect_right(ends, x + 2) - bisect.bisect_left(ends, x - 2) >= 3]


def add_text_frames(body, items, page_width):
    edges = right_edges(items)
    by_top = sorted(items, key=lambda it: it["y0"])
    tops = [it["y0"] for it in by_top]
    tallest = max((it["y1"] - it["y0"] for it in items), default=0)
    for group in group_items(items):
        x0 = min(it["x0"] for it in group)
        x1 = max(it["x1"] for it in group)
        pitch = statistics.median(b["base"] - a["base"] for a, b in zip(group, group[1:])) if len(group) > 1 else None
        line, above = line_metrics(group[0], pitch)
        top, rest = group[0]["base"] - above, group[1:]
        # по ширине: все строки, кроме последней, кончаются на одном краю — и это край колонки текста
        # (у абзаца из двух строк одна полная строка ещё ничего не доказывает)
        justified = rest and all(x1 - it["x1"] < 2 for it in group[:-1]) \
            and all(abs(it["x0"] - rest[0]["x0"]) < 2 for it in rest) \
            and (len(group) > 2 or any(abs(x1 - e) <= 2 for e in edges))
        if justified:
            align, width, x = "both", x1 - x0 + 0.5, x0  # ширина рамки — ровно как в PDF
        elif rest and all(abs((it["x0"] + it["x1"]) / 2 - (x0 + x1) / 2) < 2 for it in group) \
                and any(abs(it["x0"] - x0) > 2 for it in group):
            align, width, x = "center", x1 - x0 + 8, x0 - 4
        elif rest and all(abs(it["x1"] - x1) < 2 for it in group) and any(abs(it["x0"] - x0) > 2 for it in group):
            align, width, x = "right", x1 - x0 + 8, x0 - 8
        else:  # по левому краю: рамка до соседнего текста справа (или края страницы) — с запасом на шрифт
            room, bottom, own = page_width - 2, top + line * len(group), {id(it) for it in group}
            for other in by_top[bisect.bisect_left(tops, top - tallest):bisect.bisect_left(tops, bottom)]:
                if other["y1"] > top and other["x0"] >= x1 - 1 and id(other) not in own:
                    room = min(room, other["x0"] - 1)
            align, width, x = "left", max(room - x0, x1 - x0 + 2), x0
        first_indent = group[0]["x0"] - x0 if align in ("left", "both") else 0
        content = lines_xml([it["runs"] for it in group])
        body.append(frame_paragraph(max(x, 0), top, width, line * len(group), line, align, first_indent, content))


# ----------------------------------------------------------------------------- таблицы Word

def cell_paragraphs(cell, margin):
    """Абзацы ячейки: строки как в PDF, выравнивание по положению текста в ячейке."""
    rect = cell["rect"]
    items = sorted(cell["items"], key=lambda it: (it["base"], it["x0"]))
    lines = []  # куски на одной базовой линии — одна строка
    for item in items:
        if lines and abs(item["base"] - lines[-1][-1]["base"]) < 0.4 * item["size"]:
            lines[-1].append(item)
        else:
            lines.append([item])
    if not lines:
        return '<w:p><w:pPr><w:spacing w:before="0" w:after="0"/></w:pPr></w:p>', None, 0
    paragraphs, heights = [], []
    for k, line in enumerate(lines):
        pitch = lines[k + 1][0]["base"] - line[0]["base"] if k + 1 < len(lines) else None
        height, _ = line_metrics(line[0], pitch if pitch and pitch < 2.5 * line[0]["size"] else None)
        heights.append(height)
        x0, x1 = min(it["x0"] for it in line), max(it["x1"] for it in line)
        left, right = x0 - (rect.x0 + margin), (rect.x1 - margin) - x1
        if abs(left - right) < 3 and left > 2:
            align, indent = "center", 0
        elif right < 2.5 and left > right + 3:
            align, indent = "right", 0
        else:
            align, indent = "left", max(left, 0) if left > 2 else 0
        runs = []
        for j, item in enumerate(sorted(line, key=lambda it: it["x0"])):
            if j:
                runs.append(run_xml(" ", item["runs"][0][1]))
            runs += [run_xml(text, style) for text, style in item["runs"]]
        paragraphs.append(f'<w:p><w:pPr><w:widowControl w:val="0"/><w:snapToGrid w:val="0"/>'
                          f'<w:spacing w:before="0" w:after="0" w:line="{twips(height)}" w:lineRule="exact"/>'
                          f'<w:ind w:left="{twips(indent)}" w:right="0"/><w:jc w:val="{align}"/></w:pPr>{"".join(runs)}</w:p>')
    first, last = lines[0][0], lines[-1][0]
    top = first["base"] - line_metrics(first)[1] - rect.y0
    bottom = rect.y1 - (last["base"] - line_metrics(last)[1] + heights[-1])
    if abs(top - bottom) <= max(2, 0.2 * (top + bottom)):
        valign = "center"
    elif bottom < top:
        valign = "bottom"
    else:
        valign = "top"
        if top > 1:  # отступ сверху как в PDF
            paragraphs[0] = paragraphs[0].replace('w:before="0"', f'w:before="{twips(top)}"', 1)
    return "".join(paragraphs), valign, sum(heights) + (top if valign == "top" else 0)


def table_xml(table, horizontal, vertical, fills):
    xs, ys = table["xs"], table["ys"]
    starts = {(c["r0"], c["c0"]): c for c in table["cells"]}
    covered = {}
    for c in table["cells"]:
        for r in range(c["r0"], c["r1"]):
            for col in range(c["c0"], c["c1"]):
                covered.setdefault((r, col), c)
    offsets = [it["x0"] - c["rect"].x0 for c in table["cells"] for it in c["items"]]
    margin = max(min(min(offsets) if offsets else 2.0, 5.0), 0.5)  # поле ячейки слева/справа

    def borders(c):
        r = c["rect"]
        widths = {"top": edge_width(horizontal, r.y0, r.x0, r.x1), "bottom": edge_width(horizontal, r.y1, r.x0, r.x1),
                  "left": edge_width(vertical, r.x0, r.y0, r.y1), "right": edge_width(vertical, r.x1, r.y0, r.y1)}
        parts = []
        for side in ("top", "left", "bottom", "right"):
            w = widths[side]
            parts.append(f'<w:{side} w:val="single" w:sz="{min(max(int(round(w * 8)), 2), 96)}" w:space="0" '
                         f'w:color="000000"/>' if w else f'<w:{side} w:val="nil"/>')
        return "<w:tcBorders>" + "".join(parts) + "</w:tcBorders>"

    def shading(c):
        for rect, color in fills:
            if abs(rect & c["rect"]) > 0.8 * abs(c["rect"]) and color and min(color) < 0.97:
                hexcolor = "".join(f"{int(round(v * 255)):02X}" for v in color[:3])
                return f'<w:shd w:val="clear" w:color="auto" w:fill="{hexcolor}"/>'
        return ""

    rows = []
    for r in range(len(ys) - 1):
        height, cells_xml, col = ys[r + 1] - ys[r], [], 0
        need = 0
        while col < len(xs) - 1:
            c = covered.get((r, col))
            if c is None:  # дыра в сетке — пустая ячейка без рамок
                cells_xml.append(f'<w:tc><w:tcPr><w:tcW w:w="{twips(xs[col + 1] - xs[col])}" w:type="dxa"/>'
                                 f'<w:tcBorders><w:top w:val="nil"/><w:left w:val="nil"/><w:bottom w:val="nil"/>'
                                 f'<w:right w:val="nil"/></w:tcBorders></w:tcPr><w:p/></w:tc>')
                col += 1
                continue
            span = c["c1"] - c["c0"]
            width = twips(xs[c["c1"]] - xs[c["c0"]])
            grid_span = f'<w:gridSpan w:val="{span}"/>' if span > 1 else ""
            if (r, c["c0"]) == (c["r0"], c["c0"]) and c is starts.get((c["r0"], c["c0"])):
                content, valign, content_height = cell_paragraphs(c, margin)
                if c["r1"] - c["r0"] == 1:
                    need = max(need, content_height)
                merge = '<w:vMerge w:val="restart"/>' if c["r1"] - c["r0"] > 1 else ""
                direction = ""
                if c["items"] and all(not is_horizontal(it) for it in c["items"]):
                    direction = '<w:textDirection w:val="btLr"/>' if c["items"][0]["dir"][1] < 0 else \
                        '<w:textDirection w:val="tbRl"/>'
                valign_xml = f'<w:vAlign w:val="{valign}"/>' if valign else ""
                cells_xml.append(f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{grid_span}{merge}{borders(c)}'
                                 f'{shading(c)}{direction}{valign_xml}</w:tcPr>{content}</w:tc>')
            else:  # продолжение объединённой по вертикали ячейки
                cells_xml.append(f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{grid_span}<w:vMerge/>'
                                 f'{borders(c)}</w:tcPr><w:p/></w:tc>')
            col = c["c1"]
        rule = "exact" if need <= height + 0.5 else "atLeast"  # текст не влезает — строка подрастёт
        rows.append(f'<w:tr><w:trPr><w:cantSplit/><w:trHeight w:val="{twips(height)}" w:hRule="{rule}"/></w:trPr>'
                    f'{"".join(cells_xml)}</w:tr>')
    grid = "".join(f'<w:gridCol w:w="{twips(b - a)}"/>' for a, b in zip(xs, xs[1:]))
    return parse_xml(
        f'<w:tbl {nsdecls("w")}><w:tblPr>'
        f'<w:tblpPr w:leftFromText="0" w:rightFromText="0" w:topFromText="0" w:bottomFromText="0" '
        f'w:vertAnchor="page" w:horzAnchor="page" w:tblpX="{twips(xs[0])}" w:tblpY="{twips(ys[0])}"/>'
        f'<w:tblOverlap w:val="overlap"/><w:tblW w:w="{twips(xs[-1] - xs[0])}" w:type="dxa"/>'
        f'<w:tblLayout w:type="fixed"/><w:tblCellMar><w:top w:w="0" w:type="dxa"/>'
        f'<w:left w:w="{twips(margin)}" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/>'
        f'<w:right w:w="{twips(margin)}" w:type="dxa"/></w:tblCellMar></w:tblPr>'
        f'<w:tblGrid>{grid}</w:tblGrid>{"".join(rows)}</w:tbl>')


# ----------------------------------------------------------------------------- подложка

def background_dpi(page, dpi=BACKGROUND_DPI):
    """Разрешение подложки: для огромных страниц ниже, чтобы картинка была не больше ~25 Мпикс."""
    inches = max(page.rect.width / 72, 0.1) * max(page.rect.height / 72, 0.1)
    return max(min(dpi, (25_000_000 / inches) ** 0.5), 50)


def picture(page, dpi):
    """Картинка страницы: JPEG, если на ней есть фото (иначе файл огромный), PNG — если только линии и текст."""
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    if min(pix.samples) >= 250:  # пусто — подложка не нужна
        return None
    try:
        photos = bool(page.get_images())
    except Exception:  # noqa: BLE001
        photos = False
    return pix.tobytes("jpeg", jpg_quality=88) if photos else pix.tobytes("png")


def background(doc_pdf, pno, text_rects, table_rects):
    """Картинка страницы без перенесённого текста и без линий таблиц Word (или None, если пусто)."""
    tmp = pymupdf.open()
    tmp.insert_pdf(doc_pdf, from_page=pno, to_page=pno)
    page = tmp[0]
    if table_rects:
        for rect in table_rects:
            page.add_redact_annot(rect + (-SNAP - 1, -SNAP - 1, SNAP + 1, SNAP + 1))
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_COVERED, text=pymupdf.PDF_REDACT_TEXT_NONE)
    if text_rects:
        for rect in text_rects:
            page.add_redact_annot(rect)
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                              text=pymupdf.PDF_REDACT_TEXT_REMOVE)
    return picture(page, background_dpi(page))


def anchor_background(paragraph, image, width, height):
    """Подложка: картинка во всю страницу за текстом, привязанная к странице (не сдвигается)."""
    run = paragraph.add_run()
    run.add_picture(io.BytesIO(image), width=Pt(width), height=Pt(height))
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    doc_pr = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="0" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{extent.get("cx")}" cy="{extent.get("cy")}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/></wp:anchor>')
    anchor.append(doc_pr)
    anchor.append(parse_xml(f'<wp:cNvGraphicFramePr {nsdecls("wp")}/>'))
    anchor.append(graphic)
    drawing.remove(inline)
    drawing.append(anchor)


# ----------------------------------------------------------------------------- страница целиком

def tiny(paragraph_element):
    """Служебный абзац (подложка, разрыв раздела) высотой 1 pt — чтобы ничего не выталкивал на новую страницу."""
    p_pr = paragraph_element.get_or_add_pPr()
    for old in p_pr.findall(qn("w:spacing")) + p_pr.findall(qn("w:rPr")):
        p_pr.remove(old)
    spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>')
    sect_pr = p_pr.find(qn("w:sectPr"))
    (sect_pr.addprevious if sect_pr is not None else p_pr.append)(spacing)
    rpr = parse_xml(f'<w:rPr {nsdecls("w")}><w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr>')
    (sect_pr.addprevious if sect_pr is not None else p_pr.append)(rpr)


def normalize(doc_pdf, pno):
    """Страница без поворота и не больше 22 дюймов (предел Word). Возвращает (документ, номер страницы)."""
    page = doc_pdf[pno]
    if page.rotation:
        page.remove_rotation()  # повёрнутая страница: координаты текста и линий — как на экране
    side = max(page.rect.width, page.rect.height)
    if side <= MAX_PAGE:
        return doc_pdf, pno
    k = MAX_PAGE / side  # чертёж A1/A0: уменьшаем страницу целиком (текст остаётся текстом)
    small = pymupdf.open()
    small.new_page(width=page.rect.width * k, height=page.rect.height * k).show_pdf_page(
        pymupdf.Rect(0, 0, page.rect.width * k, page.rect.height * k), doc_pdf, pno)
    return small, 0


def layout_page(doc_pdf, pno):
    """Всё для страницы Word: (размер, подложка, элементы XML таблиц и абзацев-рамок)."""
    page = doc_pdf[pno]
    items = text_items(page)
    tables = find_word_tables(page)
    horizontal, vertical, fills = segments(page) if tables else ([], [], [])
    free, moved = [], []
    for item in items:
        center = pymupdf.Point((item["x0"] + item["x1"]) / 2, (item["y0"] + item["y1"]) / 2)
        cell = next((c for t in tables if center in t["bbox"] for c in t["cells"] if center in c["rect"]), None)
        if cell is not None:
            cell["items"].append(item)
            moved.append(item)
        elif is_horizontal(item):
            free.append(item)
            moved.append(item)
        # повёрнутый текст вне таблиц (надписи на полях рамки) остаётся на подложке
    elements = []
    for table in tables:
        elements += [table_xml(table, horizontal, vertical, fills), spacer_paragraph()]
    add_text_frames(elements, free, page.rect.width)
    image = background(doc_pdf, pno, [pymupdf.Rect(it["x0"], it["y0"], it["x1"], it["y1"]) for it in moved],
                       [t["bbox"] for t in tables])
    return page.rect, image, elements


def place_page(doc, rect, image, elements, first):
    """Новый раздел Word размером со страницу PDF, без полей: подложка и элементы на своих местах."""
    if first:
        section = doc.sections[0]
    else:
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        tiny(doc.element.body[-2])  # абзац с разрывом предыдущего раздела
    section.page_width, section.page_height = Pt(rect.width), Pt(rect.height)
    section.orientation = 1 if rect.width > rect.height else 0
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin", "header_distance", "footer_distance"):
        setattr(section, side, 0)
    holder = doc.add_paragraph()  # абзац, к которому привязана подложка
    tiny(holder._p)
    if image:
        anchor_background(holder, image, rect.width, rect.height)
    previous = holder._p
    for element in elements:
        previous.addnext(element)
        previous = element


def add_page(doc, doc_pdf, pno, first):
    """Переносит страницу pno документа doc_pdf в doc с фиксированной вёрсткой."""
    doc_pdf, pno = normalize(doc_pdf, pno)
    place_page(doc, *layout_page(doc_pdf, pno), first)


def add_image_page(doc, doc_pdf, pno, first):
    """Запасной вариант: страница целиком картинкой (если перенести её по частям не удалось)."""
    doc_pdf, pno = normalize(doc_pdf, pno)
    page = doc_pdf[pno]
    place_page(doc, page.rect, picture(page, background_dpi(page)), [], first)


def prepare(doc):
    """Настройки документа: режим Word 2013+ (современные правила размещения таблиц и рамок),
    абзацы по умолчанию без интервалов."""
    settings = doc.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = parse_xml(f'<w:compat {nsdecls("w")}/>')
        settings.append(compat)
    for old in compat.findall(qn("w:compatSetting")):
        if old.get(qn("w:name")) == "compatibilityMode":
            compat.remove(old)
    compat.append(parse_xml(f'<w:compatSetting {nsdecls("w")} w:name="compatibilityMode" '
                            f'w:uri="http://schemas.microsoft.com/office/word" w:val="15"/>'))
    spacing = doc.styles.element.find(f'{qn("w:docDefaults")}/{qn("w:pPrDefault")}/{qn("w:pPr")}/{qn("w:spacing")}')
    if spacing is not None:
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "240")
